#!/usr/bin/env python3
"""
Builds the Privileged Access Audit as a Word document on an existing branded template (for example the
PIM Phase 2 HLD). Keeps the template's page setup, styles, logo header and footer; replaces the body.

Inputs are the outputs of Invoke-PimPhase2Assessment.ps1 (the Phase2 folder) and, optionally, the Azure
review's Raw folder (for Lighthouse and coverage detail).

    python3 tools/new_audit_word_report.py --phase2 <run>/Phase2 --azure-raw <run>/Azure/Raw \
        --template "PIM Phase2 - HLD.docx" --out "Privileged Access Audit - SE.docx" \
        --tenant-name "Scottish Enterprise" --owner "Andy Kemp - PIM Subject Matter Expert"

Requires python-docx and matplotlib. The template itself is client material and is not stored in the repo.
"""
import argparse
import copy
import csv
import datetime as dt
import io
import json
import os
import re
from collections import Counter, defaultdict

import docx
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm, Emu

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

INK = "1F2933"
NAVY = "244A64"
TEAL = "2A7D7A"
LINE = "D9E1E5"
BAND = "F7F9FA"
LABEL = "EAF2F6"
RESULT_FILL = {
    "Good": ("E2EFDA", "1E6B3A"), "Achieved": ("E2EFDA", "1E6B3A"), "Strong": ("E2EFDA", "1E6B3A"), "Within": ("E2EFDA", "1E6B3A"),
    "Just-in-time (PIM)": ("E2EFDA", "1E6B3A"), "Not observed": ("E2EFDA", "1E6B3A"), "Addressed": ("E2EFDA", "1E6B3A"), "Yes": ("E2EFDA", "1E6B3A"),
    "Partial": ("FFF2CC", "8A5300"), "Partially achieved": ("FFF2CC", "8A5300"), "Partly in place": ("FFF2CC", "8A5300"), "Mixed": ("FFF2CC", "8A5300"),
    "Indicated": ("FFF2CC", "8A5300"), "Partially": ("FFF2CC", "8A5300"), "Test required": ("FFF2CC", "8A5300"), "Time-bound active": ("FFF2CC", "8A5300"),
    "Weak": ("F8D7DA", "9B1C1C"), "Not achieved": ("F8D7DA", "9B1C1C"), "Over": ("F8D7DA", "9B1C1C"), "Permanent (no PIM)": ("F8D7DA", "9B1C1C"),
    "Confirmed": ("F8D7DA", "9B1C1C"), "No": ("F8D7DA", "9B1C1C"),
    "Critical": ("7A0E0E", "FFFFFF"), "High": ("F8D7DA", "9B1C1C"), "Medium": ("FFF2CC", "8A5300"), "Low": ("EEF1F4", "44505C"),
    "Not assessed": ("EEF1F4", "5C6770"), "Open": ("EEF1F4", "5C6770"),
    "TL1": ("244A64", "FFFFFF"), "TL2": ("2A7D7A", "FFFFFF"), "TL3": ("B7D3D2", INK), "TL4": ("EEF1F4", INK),
}
TL_NAMES = {"TL1": "Strategic Control", "TL2": "Service Administration", "TL3": "Operational Administration", "TL4": "Workforce"}
HUMAN = {"CLAD", "Other admin account", "Everyday", "Guest", "Generic/Service"}

# Delivery phase for each control when it is not yet met. Generic: works for every tenant.
PHASES = [
    ("Now (0-30 days)", ["PIM-06", "ADM-04", "ADM-02", "ADM-01", "PIM-08", "PIM-05", "AZ-04", "ID-07", "ID-08"]),
    ("Next (30-90 days)", ["PIM-04", "PIM-01", "AZ-02", "AZ-01", "PIM-03", "ID-06", "ID-01", "ID-05", "ID-03", "PIM-10", "PIM-11", "PIM-07", "GOV-01", "PIM-02"]),
    ("Later (90 days +)", ["DEV-01", "ID-02", "GOV-02", "WL-01", "AZ-03", "LOG-01", "PIM-09", "ID-04", "ADM-03"]),
]


# ------------------------------------------------------------------------------------------------
# Data
# ------------------------------------------------------------------------------------------------
def read_csv(folder, name):
    if not folder:
        return []
    p = os.path.join(folder, name)
    if not os.path.exists(p) or os.path.getsize(p) < 5:
        return []
    with open(p, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def as_int(v, default=0):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return default


# ------------------------------------------------------------------------------------------------
# Low-level Word helpers (formatting copied from the template)
# ------------------------------------------------------------------------------------------------
def set_cell_fill(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_borders(cell, color=LINE, sz=4):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)


def set_cell_margins(cell, top=70, left=95, bottom=70, right=95):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tcPr.append(m)


def set_cell_width(cell, dxa):
    tcPr = cell._tc.get_or_add_tcPr()
    w = tcPr.find(qn("w:tcW"))
    if w is None:
        w = OxmlElement("w:tcW")
        tcPr.insert(0, w)
    w.set(qn("w:w"), str(int(dxa)))
    w.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trPr.append(e)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def fixed_layout(table, widths):
    tbl = table._tbl
    tblPr = tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW")
        tblPr.append(tw)
    tw.set(qn("w:w"), str(int(sum(widths))))
    tw.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for gc in list(grid):
        grid.remove(gc)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w)))
        grid.append(gc)


def cell_text(cell, text, bold=False, color=INK, size=9, italic=False, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    if align:
        p.alignment = align
    lines = str(text if text is not None else "").split("\n")
    for i, line in enumerate(lines):
        if i:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor.from_string(color)


class Builder:
    def __init__(self, template, header_text, footer_text):
        self.doc = docx.Document(template)
        self._clear_body()
        self._drop_comment_parts()
        self._set_header_footer(header_text, footer_text)
        sec = self.doc.sections[0]
        self.text_width = sec.page_width - sec.left_margin - sec.right_margin  # EMU
        self.text_dxa = int(self.text_width / 635)
        self._fig = 0

    # -- template plumbing -----------------------------------------------------------------------
    def _clear_body(self):
        body = self.doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)

    def _drop_comment_parts(self):
        part = self.doc.part
        for rid, rel in list(part.rels.items()):
            if re.search(r"/(comments|commentsExtended|commentsIds|commentsExtensible|people)$", rel.reltype):
                part.drop_rel(rid)

    def _set_header_footer(self, header_text, footer_text):
        for sec in self.doc.sections:
            for hf in (sec.header, sec.first_page_header, sec.even_page_header):
                for t in hf._element.iter(qn("w:t")):
                    if t.text and "High Level Design" in t.text:
                        t.text = header_text
            for hf in (sec.footer, sec.first_page_footer, sec.even_page_footer):
                for t in hf._element.iter(qn("w:t")):
                    if t.text and "working draft" in t.text.lower():
                        t.text = footer_text

    def set_properties(self, title, subject, author, keywords, description):
        cp = self.doc.core_properties
        cp.title, cp.subject, cp.author, cp.keywords, cp.comments = title, subject, author, keywords, description
        cp.last_modified_by = author
        cp.revision = 1
        cp.created = dt.datetime.now()
        cp.modified = dt.datetime.now()
        # Ask Word to refresh the table of contents when the document is opened.
        settings = self.doc.settings.element
        if settings.find(qn("w:updateFields")) is None:
            e = OxmlElement("w:updateFields")
            e.set(qn("w:val"), "true")
            # Schema order: updateFields comes before hdrShapeDefaults / footnotePr / endnotePr / compat ...
            after = ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema",
                     "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
                     "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags",
                     "decimalSymbol", "listSeparator")
            anchor = next((c for c in settings if c.tag.split("}")[-1] in after), None)
            if anchor is not None:
                anchor.addprevious(e)
            else:
                settings.append(e)

    # -- content -----------------------------------------------------------------------------------
    def h(self, text, level=1):
        return self.doc.add_paragraph(text, style=f"Heading {level}")

    def p(self, text="", bold_lead=None, size=None, color=None, italic=False, after=None):
        para = self.doc.add_paragraph(style="Normal")
        if bold_lead:
            r = para.add_run(bold_lead)
            r.bold = True
            if size:
                r.font.size = Pt(size)
        if text:
            r = para.add_run(text)
            r.italic = italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        return para

    def bullet(self, text, bold_lead=None):
        para = self.doc.add_paragraph(style="List Bullet")
        if bold_lead:
            para.add_run(bold_lead).bold = True
        para.add_run(text)
        return para

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def table(self, headers, rows, widths, badge_cols=(), size=8.5, header_fill=NAVY, band=True):
        """widths are relative; scaled to the text width."""
        total = float(sum(widths))
        dxa = [w / total * self.text_dxa for w in widths]
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        fixed_layout(t, dxa)
        hdr = t.rows[0]
        set_repeat_header(hdr)
        no_split(hdr)
        for i, htext in enumerate(headers):
            c = hdr.cells[i]
            set_cell_width(c, dxa[i])
            set_cell_fill(c, header_fill)
            set_cell_borders(c, header_fill)
            set_cell_margins(c)
            cell_text(c, htext, bold=True, color="FFFFFF", size=size)
        for ri, row in enumerate(rows):
            tr = t.rows[ri + 1]
            no_split(tr)
            for ci, val in enumerate(row):
                c = tr.cells[ci]
                set_cell_width(c, dxa[ci])
                set_cell_borders(c)
                set_cell_margins(c)
                sval = "" if val is None else str(val)
                if ci in badge_cols and sval in RESULT_FILL:
                    fill, fg = RESULT_FILL[sval]
                    set_cell_fill(c, fill)
                    cell_text(c, sval, bold=True, color=fg, size=size)
                else:
                    if band and ri % 2 == 1:
                        set_cell_fill(c, BAND)
                    cell_text(c, sval, size=size)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def kv_table(self, pairs, widths=(2232, 6912)):
        total = float(sum(widths))
        dxa = [w / total * self.text_dxa for w in widths]
        t = self.doc.add_table(rows=len(pairs), cols=2)
        fixed_layout(t, dxa)
        for i, (k, v) in enumerate(pairs):
            a, b = t.rows[i].cells
            for c, w in ((a, dxa[0]), (b, dxa[1])):
                set_cell_width(c, w)
                set_cell_borders(c)
                set_cell_margins(c, 85, 95, 85, 95)
            set_cell_fill(a, LABEL)
            cell_text(a, k, bold=True, size=9)
            cell_text(b, v, size=9)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def callout(self, title, lines, kind="teal"):
        fill, border, tcolor = {"teal": ("E8F4F2", TEAL, TEAL), "amber": ("FFF2CC", "F4B183", "C65911"), "red": ("FBE4E4", "E8A3A3", "9B1C1C"), "green": ("E2EFDA", "A9D18E", "1E6B3A")}[kind]
        t = self.doc.add_table(rows=1, cols=1)
        fixed_layout(t, [self.text_dxa])
        c = t.rows[0].cells[0]
        set_cell_width(c, self.text_dxa)
        set_cell_fill(c, fill)
        set_cell_borders(c, border, 10)
        set_cell_margins(c, 120, 150, 120, 150)
        cell_text(c, title, bold=True, color=tcolor, size=10)
        for line in lines:
            para = c.add_paragraph()
            para.paragraph_format.space_after = Pt(2)
            if isinstance(line, tuple):
                r = para.add_run(line[0])
                r.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor.from_string(INK)
                line = line[1]
            r = para.add_run(line)
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor.from_string(INK)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def toc(self):
        para = self.doc.add_paragraph()
        run = para.add_run()
        for tag, text in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), (None, "Right-click and choose Update Field to build the table of contents."), ("end", None)):
            if tag:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), tag)
                run._r.append(fc)
            elif text.startswith("TOC"):
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = text
                run._r.append(it)
            else:
                t = OxmlElement("w:t")
                t.text = text
                run._r.append(t)

    def figure(self, fig, caption, width_frac=1.0):
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=200, bbox_inches="tight")
        plt.close(fig)
        buf.seek(0)
        self.doc.add_paragraph().add_run().add_picture(buf, width=Emu(int(self.text_width * width_frac)))
        self._fig += 1
        self.p(f"Figure {self._fig}: {caption}", size=8.5, color="5C6770", italic=True)

    def cover(self, title_lines, subtitle, meta, status_title, status_lines):
        for i, line in enumerate(title_lines):
            para = self.doc.add_paragraph(style="Title")
            if i == 0:
                para.paragraph_format.space_after = Pt(28)
            r = para.add_run(line)
            r.font.name = "Arial"
            r.font.size = Pt(25)
        para = self.doc.add_paragraph()
        para.paragraph_format.space_after = Pt(18)
        r = para.add_run(subtitle)
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor.from_string(TEAL)
        bar = self.doc.add_table(rows=1, cols=1)
        fixed_layout(bar, [2880])
        c = bar.rows[0].cells[0]
        set_cell_width(c, 2880)
        set_cell_fill(c, TEAL)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(14)
        self.kv_table(meta)
        self.callout(status_title, status_lines, "amber")


TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]


def _reorder(el, order):
    kids = list(el)
    rank = {n: i for i, n in enumerate(order)}
    kids.sort(key=lambda c: rank.get(c.tag.split("}")[-1], len(order)))
    seen = set()
    for c in kids:
        el.remove(c)
    for c in kids:
        name = c.tag.split("}")[-1]
        if name in seen and name in ("tcBorders", "shd", "tcMar", "tblLayout", "tcW"):
            continue  # keep the first of any duplicate
        seen.add(name)
        el.append(c)


def normalise_order(document):
    """Put table property elements in schema order so Word opens the file without repair."""
    body = document.element.body
    for tcPr in body.iter(qn("w:tcPr")):
        _reorder(tcPr, TCPR_ORDER)
    for tblPr in body.iter(qn("w:tblPr")):
        _reorder(tblPr, TBLPR_ORDER)


# ------------------------------------------------------------------------------------------------
# Charts
# ------------------------------------------------------------------------------------------------
def chart_pim_coverage(paths):
    rows = []
    for plane in ("Entra", "Azure"):
        for tl in ("TL1", "TL2", "TL3"):
            rs = [p for p in paths if p["Plane"] == plane and p["TrustLevel"] == tl and p.get("AzureFamily", "") != "Reader"]
            if not rs:
                continue
            j = sum(1 for p in rs if p["EffectiveState"] == "Eligible")
            perm = sum(1 for p in rs if p["Standing"] == "True")
            tb = max(0, len(rs) - j - perm)
            rows.append((f"{plane} {tl}", j, tb, perm, len(rs)))
    if not rows:
        return None
    fig, ax = plt.subplots(figsize=(7.2, 0.45 * len(rows) + 0.9))
    labels = [r[0] for r in rows][::-1]
    data = rows[::-1]
    left = [0] * len(data)
    for idx, (name, color) in enumerate((("Just-in-time (PIM)", "#2A7D7A"), ("Time-bound active", "#E0A243"), ("Permanent", "#B42318"))):
        vals = [100.0 * r[1 + idx] / r[4] for r in data]
        ax.barh(labels, vals, left=left, color=color, label=name, height=0.6)
        for k, v in enumerate(vals):
            if v >= 7:
                ax.text(left[k] + v / 2, k, f"{v:.0f}%", va="center", ha="center", fontsize=7.5, color="white")
        left = [a + b for a, b in zip(left, vals)]
    for k, r in enumerate(data):
        ax.text(101, k, f"{r[4]} routes", va="center", fontsize=7.5, color="#5C6770")
    ax.set_xlim(0, 115)
    ax.set_xticks([0, 25, 50, 75, 100])
    ax.set_xticklabels(["0%", "25%", "50%", "75%", "100%"], fontsize=8)
    ax.tick_params(axis="y", labelsize=8.5)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.legend(loc="upper center", bbox_to_anchor=(0.45, -0.12), ncol=3, frameon=False, fontsize=8)
    return fig


def chart_identities_by_tl(identities):
    cats = [("Admin account (CLAD)", {"CLAD", "Other admin account"}, "#244A64"), ("Everyday or guest", {"Everyday", "Guest"}, "#B42318"), ("Shared / service", {"Generic/Service"}, "#E0A243")]
    tls = [t for t in ("TL1", "TL2", "TL3") if any(i["HighestTrustLevel"] == t for i in identities)]
    if not tls:
        return None
    fig, ax = plt.subplots(figsize=(7.2, 2.4))
    left = [0] * len(tls)
    labels = [f"{t} {TL_NAMES[t]}" for t in tls][::-1]
    for name, types, color in cats:
        vals = [sum(1 for i in identities if i["HighestTrustLevel"] == t and i["AccountType"] in types) for t in tls][::-1]
        ax.barh(labels, vals, left=left, color=color, label=name, height=0.6)
        for k, v in enumerate(vals):
            if v >= 4:
                ax.text(left[k] + v / 2, k, str(v), va="center", ha="center", fontsize=8, color="white")
        left = [a + b for a, b in zip(left, vals)]
    ax.tick_params(axis="y", labelsize=8.5)
    ax.tick_params(axis="x", labelsize=8)
    ax.set_xlabel("Identities (each counted at its highest trust level)", fontsize=8)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.legend(loc="upper center", bbox_to_anchor=(0.45, -0.28), ncol=3, frameon=False, fontsize=8)
    return fig


def chart_oversubscribed(over):
    rows = [o for o in over if o["Status"] == "Over" and o["Plane"] == "Entra" and o["TrustLevel"] == "TL1" and not o["Role"].startswith("Any ")]
    rows = sorted(rows, key=lambda o: -as_int(o["People"]))[:10]
    if not rows:
        return None
    fig, ax = plt.subplots(figsize=(7.2, 0.38 * len(rows) + 0.9))
    labels = [o["Role"] for o in rows][::-1]
    people = [as_int(o["People"]) for o in rows][::-1]
    guide = [as_int(o["Guide"]) for o in rows][::-1]
    ax.barh(labels, people, color="#B42318", height=0.55, label="People who can reach the role")
    ax.scatter(guide, range(len(guide)), marker="|", s=260, color="#1F2933", label="Guide maximum", zorder=3)
    for k, v in enumerate(people):
        ax.text(v + 0.5, k, str(v), va="center", fontsize=8)
    ax.tick_params(axis="y", labelsize=8)
    ax.tick_params(axis="x", labelsize=8)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.legend(loc="lower right", frameon=False, fontsize=8)
    return fig


# ------------------------------------------------------------------------------------------------
# Report
# ------------------------------------------------------------------------------------------------
def build(args):
    ph = args.phase2
    metrics = (read_csv(ph, "P2_Metrics.csv") or [{}])[0]
    controls = read_csv(ph, "P2_Controls.csv")
    caf = read_csv(ph, "P2_NcscCafBenchmark.csv")
    coverage = read_csv(ph, "P2_PimCoverage.csv")
    over = read_csv(ph, "P2_OversubscribedRoles.csv")
    hld = read_csv(ph, "P2_HldFindings.csv")
    maturity = read_csv(ph, "P2_Maturity.csv")
    identities = [i for i in read_csv(ph, "P2_PrivilegedIdentities.csv")]
    paths = [p for p in read_csv(ph, "P2_AccessPaths.csv") if p["AccountType"] in HUMAN]
    approvers = read_csv(ph, "P2_ApproverGroups.csv")
    bundles = read_csv(ph, "P2_GroupBundles.csv")
    catalogue = read_csv(ph, "P2_RoleCatalogue.csv")
    exceptions = read_csv(ph, "P2_ExceptionCandidates.csv")
    lighthouse = read_csv(args.azure_raw, "AzureLighthouseDelegations.csv")
    gaps = read_csv(args.azure_raw, "AzureCoverageGaps.csv")

    tenant = args.tenant_name or metrics.get("TenantName") or "Tenant"
    short = metrics.get("Tenant") or tenant
    today = dt.date.today().strftime("%d %B %Y").lstrip("0")
    has_azure = bool(metrics.get("AzureRun"))

    b = Builder(args.template, f"Privileged Access Audit - {tenant}  |  v{args.version}", f"{args.classification} - privileged access audit")
    b.set_properties(f"Privileged Access Audit - {tenant}", "Entra ID and Azure privileged access measured against the Trust Level model and NCSC CAF", args.owner.split(" - ")[0],
                     "PIM, privileged access, trust levels, NCSC CAF, Entra, Azure", f"Draft v{args.version}")

    human_ids = [i for i in identities if i["AccountType"] in HUMAN]
    good = [c for c in controls if c["Result"] == "Good"]
    weak = [c for c in controls if c["Result"] == "Weak"]
    part = [c for c in controls if c["Result"] == "Partial"]
    sev_rank = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3, "": 4}
    attention = sorted(weak, key=lambda c: sev_rank.get(c["Severity"], 4)) + sorted(part, key=lambda c: sev_rank.get(c["Severity"], 4))
    assessed = [c for c in controls if c["Result"] != "Not assessed"]

    # ---- Cover --------------------------------------------------------------------------------
    b.cover(["Privileged Access Audit", tenant],
            "Entra ID and Azure privileged access measured against the Trust Level model and NCSC CAF",
            [("Document owner", args.owner), ("Tenant", f"{tenant} ({short})"), ("Version", args.version),
             ("Status", "Draft for review"), ("Evidence date", f"Entra run {metrics.get('EntraRun', '')}; Azure run {metrics.get('AzureRun', 'not included')}"),
             ("Date", today), ("Audience", args.audience), ("Classification", args.classification)],
            "Draft status",
            ["This audit reports the privileged access position of the tenant at the evidence date. It is produced from configuration exported from Entra ID and Azure and is an input to PIM Phase 2; it does not approve any change.",
             "Ratings against the NCSC Cyber Assessment Framework are indicative: they rely on configuration evidence only."])
    b.page_break()
    para = b.doc.add_paragraph()
    r = para.add_run("Table of Contents")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    b.toc()
    b.page_break()

    # ---- 1 Purpose ---------------------------------------------------------------------------
    b.h("1. Document Purpose and Scope", 1)
    b.h("1.1 Purpose", 2)
    b.p(f"This document records the current privileged access position of the {tenant} tenant. It shows who can obtain administrative rights in Entra ID and Azure, how they obtain them, where Privileged Identity Management (PIM) is used and where it is not, and how the controls compare with the target Trust Level model and the NCSC Cyber Assessment Framework.")
    b.p("It is the evidence baseline for PIM Phase 2. The same audit is run for each partner tenant so the results can be compared and tracked over time.")
    b.h("1.2 Scope", 2)
    b.bullet("Microsoft Entra ID directory roles: direct assignments, PIM eligible and active assignments, role-assignable groups, PIM for Groups and nested groups.", "Entra ID: ")
    b.bullet("Every Azure role assignment from management groups down to individual resources, including PIM for Azure resources, group expansion and Azure Lighthouse delegations." if has_azure else "Azure was not included in this run.", "Azure: ")
    b.bullet("PIM role settings, approvers, Conditional Access policies, authentication method registration, licensing and access reviews.", "Controls: ")
    b.bullet("Non-human identities are inventoried but are outside the scope of PIM Phase 2.", "Out of scope: ")
    b.h("1.3 How to read this document", 2)
    b.bullet("Section 2 is the summary for decision makers.")
    b.bullet("Sections 3 and 4 explain the Trust Level model and show the tenant's admin population against it.")
    b.bullet("Sections 5 and 6 set out what is working well and what needs attention, each with the reason it matters and the recommended action.")
    b.bullet("Sections 7 to 9 give the supporting detail: PIM usage, over-subscribed roles and specific risk areas.")
    b.bullet("Section 10 benchmarks the tenant against NCSC guidance; Section 11 maps the results to the PIM Phase 2 HLD; Section 12 is the recommended plan.")
    b.bullet("Individuals are not named in this document. Named detail is in the evidence files listed in Appendix C.")

    # ---- 2 Executive summary -------------------------------------------------------------------
    b.page_break()
    b.h("2. Executive Summary", 1)
    b.h("2.1 Overall position", 2)
    gap_texts = [c["Gap"] for c in attention if c["Result"] == "Weak" and c.get("Gap")][:5]
    good_titles = [c["Control"] for c in good][:4]
    b.p(f"Of {len(assessed)} controls assessed, {len(good)} are in place, {len(part)} are partly in place and {len(weak)} are weak. "
        f"PIM is established for Entra ID: {metrics.get('EligibleSharePercent', '')}% of all privileged routes need activation, and nobody holds Global Administrator permanently. "
        "However, the way PIM has been built means it does not yet give the control it should: approval is not independent, TL1 roles are bundled with routine roles in the same groups, and too many people can reach the highest-impact roles.")
    caf_text = "; ".join(f"{c['Outcome']} {c['Title'].lower()}: {c['Rating'].lower()}" for c in caf)
    b.callout("Headline", [
        ("Working well: ", "; ".join(t[0].lower() + t[1:] if not re.match(r"^(TL\d|Global|Azure|Entra|Conditional|PIM)", t) else t for t in good_titles) + "."),
        ("Most important gaps: ", "; ".join(gap_texts) + "."),
        ("NCSC CAF B2 (indicative): ", caf_text + "."),
    ], "teal")
    b.h("2.2 Key measures", 2)
    km = [
        ("People with admin access (Entra or Azure)", metrics.get("PeopleWithPrivilege", ""), ""),
        ("People who can reach TL1 Strategic Control", metrics.get("PeopleWithTL1", ""), ""),
        ("People who can become Global Administrator", metrics.get("PeopleReachingGlobalAdmin", ""), "Fewer than 5 (Microsoft)"),
        ("People who can approve their own elevation", metrics.get("SelfApprovalPeople", ""), "None"),
        ("Groups that bundle TL1 roles with other roles", metrics.get("TL1InBundles", ""), "None"),
        ("Permanent TL1 routes held by people", metrics.get("StandingTL1Paths", ""), "None (break-glass only)"),
        ("Everyday or guest accounts holding admin rights", as_int(metrics.get("EverydayIdentitiesWithPrivilege")) + as_int(metrics.get("GuestIdentitiesWithPrivilege")), "None"),
        ("TL1 identities without a phishing-resistant method", metrics.get("TL1WithoutPhishingResistant", ""), "None (FIDO2 for TL1)"),
        ("Shared or service user accounts with admin rights", metrics.get("GenericAccounts", ""), "None, or recorded exceptions"),
    ]
    if has_azure:
        km += [("Permanent TL1 Azure routes at subscription scope or above", metrics.get("AzureStandingTL1AtSubOrAbove", ""), "None"),
               ("Azure Lighthouse authorisations to other organisations", metrics.get("LighthouseAuthorisations", ""), "Eligible, reviewed, contract-bound"),
               ("Deleted identities still holding Azure roles", metrics.get("DeletedIdentitiesWithAccess", ""), "None")]
    b.table(["Measure", "Current", "Target"], km, [6, 1.4, 2.6])
    b.h("2.3 Priority actions", 2)
    for c in [c for c in attention if c["Result"] == "Weak"][:5]:
        b.bullet(c["RecommendedAction"], f"{c['Control']}. ")

    # ---- 3 Trust level model -----------------------------------------------------------------
    b.h("3. The Trust Level Model", 1)
    b.p("The audit measures the tenant against the Trust Level model used for PIM Phase 2. Identities and roles are grouped by the impact if they were compromised, not by the system they manage, and the controls required rise with that impact. TL1 is the cloud-era counterpart of Tier 0: it protects the tenant control plane.")
    b.table(["Level", "Name", "Meaning", "Examples", "Expected controls"], [
        ["TL1", "Strategic Control", "Controls the tenant control plane; compromise could mean losing control of the whole tenant", "Global Administrator, Privileged Role Administrator, Conditional Access Administrator, Authentication Policy Administrator; Azure Owner or User Access Administrator at subscription scope or above", "Eligible only, direct or single-purpose group; FIDO2 security key; privileged access workstation; independent approval; justification; 2-hour activation; continuous monitoring; quarterly review"],
        ["TL2", "Service Administration", "Administers a critical service within its own boundary", "Exchange, SharePoint, Teams, Intune, Power Platform administrators; Azure Contributor at subscription", "Eligible; passkey or FIDO2; approval by impact; 4-hour activation; six-monthly review"],
        ["TL3", "Operational Administration", "Supports users and routine administration", "Helpdesk, user and group administration, resource-level contributor", "Eligible; phishing-resistant or strong MFA; no or light approval; 8-hour activation; annual review"],
        ["TL4", "Workforce", "Everyday identities with no privileged role", "Standard users; low-impact read-only roles", "Workforce MFA or passkeys, compliant devices, Zero Trust baseline"],
    ], [0.8, 1.6, 2.6, 3.2, 3.6], badge_cols=(0,))
    b.p("Azure roles are classified by role and scope together (HLD 7.11.2): the same role is a higher level at management group or subscription scope than on a single resource, and write access to scopes holding identity, key management, security, logging or automation services is TL1 at any scope.", size=9.5)

    # ---- 4 Position by trust level -------------------------------------------------------------
    b.h("4. Current Position by Trust Level", 1)
    b.p("Each admin identity is counted once, at the highest trust level it can reach. The target is that TL1 and TL2 are held only by dedicated admin accounts, with phishing-resistant authentication and no permanent access.")
    rows = []
    for tl in ("TL1", "TL2", "TL3"):
        g = [i for i in human_ids if i["HighestTrustLevel"] == tl]
        if not g:
            continue
        known = [i for i in g if i["PhishingResistant"] in ("True", "False")]
        pr = sum(1 for i in known if i["PhishingResistant"] == "True")
        rows.append([tl, TL_NAMES[tl], len(g), sum(1 for i in g if i["AccountType"] in ("CLAD", "Other admin account")),
                     sum(1 for i in g if i["AccountType"] in ("Everyday", "Guest")), sum(1 for i in g if i["AccountType"] == "Generic/Service"),
                     f"{round(100 * pr / len(known))}%" if known else "n/a", sum(1 for i in g if as_int(i["StandingPaths"]) > 0)])
    b.table(["Level", "Name", "Identities", "Admin accounts", "Everyday or guest", "Shared / service", "Phishing-resistant MFA", "Any permanent access"], rows, [0.8, 2.2, 1.1, 1.2, 1.3, 1.2, 1.4, 1.3], badge_cols=(0,))
    fig = chart_identities_by_tl(human_ids)
    if fig:
        b.figure(fig, "Admin identities by highest trust level and account type")
    orgs = Counter(i.get("PersonOrganisation") or "Not known" for i in human_ids)
    b.p("Admin identities by organisation: " + "; ".join(f"{k} {v}" for k, v in orgs.most_common()) + ".", size=9.5)

    # ---- 5 Working well ----------------------------------------------------------------------
    b.h("5. What Is Working Well", 1)
    b.p("These controls already meet the expected standard and should be kept as the programme changes other things.")
    b.table(["Control", "Evidence"], [[c["Control"], c["Finding"]] for c in good], [4, 6])
    if part:
        b.p(f"A further {len(part)} controls are partly in place; they are covered in Section 6.", size=9.5)

    # ---- 6 Findings --------------------------------------------------------------------------
    b.h("6. Findings Requiring Attention", 1)
    b.p("Weak controls are listed first, by severity, followed by controls that are partly in place. Each finding shows what was found, why it matters, what good looks like and the recommended action.")
    for n, c in enumerate(attention, 1):
        b.h(f"6.{n} {c['Control']}", 2)
        b.table(["Result", "Severity", "NCSC CAF", "NCSC principle", "Reference"],
                [[c["Result"], c["Severity"] or "-", c["CafOutcome"], c["NcscPrinciple"].replace("SSA ", "Secure admin "), c["Id"]]],
                [1.2, 1.1, 1.3, 1.8, 1.0], badge_cols=(0, 1), band=False)
        b.p(c["Finding"], bold_lead="Finding: ")
        if c["WhyItMatters"]:
            b.p(c["WhyItMatters"], bold_lead="Why it matters: ")
        if c["WhatGoodLooksLike"]:
            b.p(c["WhatGoodLooksLike"], bold_lead="What good looks like: ")
        if c["RecommendedAction"]:
            b.p(c["RecommendedAction"], bold_lead="Recommended action: ")
        refs = "; ".join(x for x in (c.get("Reference"), ("Evidence: " + c["Evidence"]) if c.get("Evidence") else "") if x)
        if refs:
            b.p(refs, size=8.5, color="5C6770")

    # ---- 7 PIM coverage ----------------------------------------------------------------------
    b.page_break()
    b.h("7. Where PIM Is Used, and Where It Is Not", 1)
    b.p("Every route a person has to an admin role is classed as just-in-time (eligible in PIM and activated when needed), time-bound (active with an end date, usually a current activation) or permanent. A route through a group is just-in-time when either the group membership or the group's role is eligible. Read-only Azure roles are excluded.")
    fig = chart_pim_coverage(paths)
    if fig:
        b.figure(fig, "Share of privileged routes by activation model, per plane and trust level")
    ent = [c for c in coverage if c["Plane"] == "Entra" and c["TrustLevel"] in ("TL1", "TL2")]
    b.h("7.1 Entra ID roles (TL1 and TL2)", 2)
    b.table(["Level", "Role", "PIM status", "People", "Just-in-time", "Permanent"],
            [[c["TrustLevel"], c["Role"], c["PimStatus"], c["People"], c["JustInTime"], c["Permanent"]] for c in ent],
            [0.8, 4, 2.2, 0.9, 1.2, 1.1], badge_cols=(0, 2))
    if has_azure:
        az = [c for c in coverage if c["Plane"] == "Azure"]
        b.h("7.2 Azure roles (TL1 and TL2, by scope)", 2)
        az_sorted = sorted(az, key=lambda c: (c["TrustLevel"], -as_int(c["People"])))[:40]
        b.table(["Level", "Role", "Scope", "PIM status", "People", "Just-in-time", "Permanent"],
                [[c["TrustLevel"], c["Role"], c["ScopeLevel"], c["PimStatus"], c["People"], c["JustInTime"], c["Permanent"]] for c in az_sorted],
                [0.8, 3.2, 1.6, 2.0, 0.9, 1.2, 1.1], badge_cols=(0, 3))
        if len(az) > 40:
            b.p(f"Showing 40 of {len(az)} role and scope combinations; the full list is in P2_PimCoverage.csv.", size=8.5, color="5C6770")

    # ---- 8 Over-subscribed -------------------------------------------------------------------
    b.h("8. Over-Subscribed Roles", 1)
    b.p("How many people can reach each high-impact role, against a guide number. Global Administrator uses Microsoft's guidance of fewer than five; Azure subscription owners use Microsoft's guidance of no more than three. Other roles use the assessment guide in the role catalogue (TL1: 5 people, TL2: 10), which the partners can change.")
    fig = chart_oversubscribed(over)
    if fig:
        b.figure(fig, "TL1 Entra roles above their guide: people who can reach the role against the guide maximum")
    over_rows = [o for o in over if o["Status"] == "Over" and o["TrustLevel"] == "TL1"]
    b.table(["Role", "Scope", "People", "Guide", "x Guide", "Source of guide"],
            [[o["Role"], o["Scope"], o["People"], o["Guide"], o["TimesGuide"], o["GuideSource"]] for o in over_rows],
            [3.4, 2.6, 0.9, 0.9, 1.0, 2.9])
    n_tl2 = sum(1 for o in over if o["Status"] == "Over" and o["TrustLevel"] == "TL2")
    if n_tl2:
        b.p(f"A further {n_tl2} TL2 roles exceed their guide; they are listed in P2_OversubscribedRoles.csv.", size=9.5)

    # ---- 9 Risk areas ------------------------------------------------------------------------
    b.h("9. Specific Risk Areas", 1)
    b.h("9.1 Approval and self-approval", 2)
    b.p("PIM only prevents the same account from approving its own request. Where a person's everyday account is an approver for a role their admin account can request, approval gives no independent assurance. The HLD (7.6.3) requires approver groups to contain admin identities only, separate from the population they approve.")
    b.table(["Approver group", "Highest level", "Roles approved", "Members", "Admin accounts", "Everyday or guest", "Approvers also eligible", "Can self-approve", "Meets HLD 7.6.3"],
            [[a["ApproverGroup"], a["HighestTrustLevel"], a["RolesApproved"], a["Members"], a["CladMembers"], a["EverydayOrGuestMembers"], a["ApproversAlsoEligible"], a["SelfApprovalPeople"], "Yes" if a["MeetsHld763"] == "True" else "No"] for a in approvers],
            [2.4, 1.0, 1.0, 0.9, 1.0, 1.1, 1.2, 1.1, 1.1], badge_cols=(1, 8))
    b.h("9.2 Role bundling through groups", 2)
    b.p("Groups that grant several roles at once mean one activation gives far more than the task needs. Where a TL1 role sits in such a group, Global Administrator and similar roles cannot be governed on their own.")
    brows = [[g["Group"], g["HighestTrustLevel"], g["Roles"], len([x for x in g["TL1Roles"].split(";") if x.strip()]), g["People"], g["Recommendation"]] for g in bundles if as_int(g["Roles"]) > 1][:15]
    b.table(["Group", "Highest level", "Roles", "TL1 roles", "People", "Recommendation"], brows, [3.0, 1.0, 0.8, 0.9, 0.8, 4.2], badge_cols=(1,))
    if has_azure:
        b.h("9.3 Azure", 2)
        subs = [o for o in over if o["Plane"] == "Azure"]
        if subs:
            b.p("People with Owner, User Access Administrator or RBAC Administrator at subscription scope (Microsoft guidance: no more than three per subscription):")
            b.table(["Subscription", "People", "Status"], [[o["Scope"].replace("Subscription: ", ""), o["People"], o["Status"]] for o in subs], [6, 1.2, 1.4], badge_cols=(2,))
        if lighthouse:
            b.p("Azure Lighthouse lets people in another organisation's tenant act in these subscriptions. Their accounts, MFA and devices are controlled by that organisation and their access does not appear in this tenant's PIM.")
            by = defaultdict(lambda: [0, 0, 0, set()])
            for l in lighthouse:
                k = l.get("ManagedByTenant") or l.get("ManagedByTenantId")
                write = l.get("PrivilegeTier") in ("Critical", "High") or ("reader" not in (l.get("RoleName") or "").lower() and l.get("PrivilegeTier") != "Low")
                by[k][0] += 1
                if write:
                    by[k][1] += 1
                    if (l.get("AuthorizationType") or "") != "Eligible":
                        by[k][2] += 1
                by[k][3].add(l.get("SubscriptionName"))
            b.table(["Managing organisation", "Authorisations", "Write-capable", "Permanent write-capable", "Subscriptions"],
                    [[k, v[0], v[1], v[2], len(v[3])] for k, v in sorted(by.items(), key=lambda kv: -kv[1][0])], [3.4, 1.4, 1.4, 1.8, 1.4])
    clad = read_csv(ph, "P2_CladProvisioning.csv")
    if clad:
        b.h("9.4 Admin accounts (CLAD) for everyone doing admin work", 2)
        mv = [c for c in clad if "existing CLAD" in c["Action"]]
        cr = [c for c in clad if "create clad" in c["Action"].lower()]
        conf = [c for c in mv if c["ExistingCladMatch"] != "Same person"]
        b.p(f"Every person who does admin work from an everyday or guest account, including PIM approvers, should have a dedicated admin account (CLAD) and use only that for privileged access. {len(clad)} people are in this position: {len(mv)} already have a CLAD (their rights or approver role need moving to it) and {len(cr)} need a CLAD created.")
        b.p("New CLADs follow the existing naming pattern: clad_<surname><first initial> (adding letters of the first name where that name is already taken), with the display name \"(Cloud Admin) First Last\". The proposed names, the rights to move and the approver changes are in P2_CladProvisioning.csv.")
        rows = []
        for tl in ("TL1", "TL2", "TL3"):
            g = [c for c in clad if c["HighestTrustLevel"] == tl]
            if g:
                rows.append([tl, len(g), sum(1 for c in g if "existing CLAD" in c["Action"]), sum(1 for c in g if "create clad" in c["Action"].lower()), sum(1 for c in g if c["IsApprover"] == "True"), sum(as_int(c["RoutesToMove"]) for c in g)])
        appr_only = [c for c in clad if not c["HighestTrustLevel"]]
        if appr_only:
            rows.append(["Approver only", len(appr_only), sum(1 for c in appr_only if "existing CLAD" in c["Action"]), sum(1 for c in appr_only if "create clad" in c["Action"].lower()), len(appr_only), 0])
        b.table(["Highest level", "People", "CLAD exists", "CLAD to create", "Approvers", "Routes to move"], rows, [1.6, 1.0, 2.0, 1.6, 1.2, 1.4], badge_cols=(0,))
        orgs = Counter(c["Organisation"] or "Not known" for c in clad)
        b.p("By organisation: " + "; ".join(f"{k} {v}" for k, v in orgs.most_common()) + ".", size=9.5)
        if conf:
            b.p(f"{len(conf)} existing CLADs were matched by naming pattern only (for example where an older CLAD shortens the surname) and should be confirmed before rights are moved.", size=9.5)
    plan = read_csv(ph, "P2_AccessReviewPlan.csv")
    existing_reviews = read_csv(os.path.join(os.path.dirname(ph.rstrip("/\\")), "Raw"), "AccessReviews.beta.csv") if ph else []
    if plan:
        b.h("9.5 Access reviews", 2)
        b.p(f"Access reviews keep privileged access correct after the remediation: they confirm each assignment is still needed and remove it when it is not. {len(existing_reviews) if existing_reviews else metrics.get('AccessReviewsConfigured', 0)} access reviews are configured today and none is confirmed to cover the privileged roles and PIM groups as a whole. The proposed plan gives every privileged role, PIM group and Azure subscription a review on its trust-level cycle.")
        rows = []
        for tl in ("TL1", "TL2", "TL3"):
            g = [r for r in plan if r["TrustLevel"] == tl]
            if g:
                kinds = Counter(r["ReviewType"].split(" (")[0].replace("PIM access review of ", "").replace("Access review of ", "") for r in g)
                rows.append([tl, len(g), g[0]["Frequency"], g[0]["DurationDays"] + " days", g[0]["Reviewers"], "; ".join(f"{v} {k}" for k, v in kinds.most_common())])
        b.table(["Level", "Reviews", "Frequency", "Duration", "Reviewers", "What is reviewed"], rows, [0.8, 1.1, 1.3, 1.0, 3.3, 3.5], badge_cols=(0,))
        b.p("All reviews: apply results automatically, remove access that is not confirmed, require a justification and show reviewer recommendations. Reviews of Azure resource roles need PIM for Azure resources enabled on the subscription. Access reviews of privileged roles need Entra ID P2 or ID Governance licences for the people reviewed. The full plan is in P2_AccessReviewPlan.csv.", size=9.5)
        if existing_reviews:
            b.table(["Existing review", "Status", "Default decision", "Justification"], [[r.get("DisplayName", "").strip(), r.get("Status", ""), r.get("DefaultDecision", ""), r.get("JustificationRequired", "")] for r in existing_reviews], [5, 1.4, 1.6, 1.4])
    b.h("9.6 Third-party, guest and shared accounts", 2)
    cats = Counter(e["Category"] for e in exceptions)
    b.p("Accounts that will need replacing, converting or recording as exceptions under the PIM Phase 2 exception process (HLD 8.4):")
    b.table(["Category", "Count", "Required evidence for an exception"],
            [[k, v, next((e["RequiredEvidence"] for e in exceptions if e["Category"] == k), "")] for k, v in cats.most_common()], [3.2, 0.8, 6])
    generic = [i for i in human_ids if i["AccountType"] == "Generic/Service"]
    if generic:
        b.p("Shared or service user accounts holding admin rights: " + "; ".join(sorted({i["Identity"] for i in generic})) + ".", size=9.5)

    # ---- 10 NCSC -----------------------------------------------------------------------------
    b.h("10. NCSC Benchmark", 1)
    b.p("The NCSC Cyber Assessment Framework (CAF 4.0) is the benchmark for UK public sector organisations, including under GovAssure. Principle B2 covers identity and access control. The ratings below are indicative: configuration evidence alone cannot show process indicators such as joiners, movers and leavers records, review sign-offs or monitoring, so those are marked Not assessed. A single weak control at High or Critical severity gives Not achieved, mirroring the CAF approach.")
    b.table(["Outcome", "Title", "Indicative rating", "Good", "Partial", "Weak", "Controls"],
            [[c["Outcome"], c["Title"], c["Rating"], c["Good"], c["Partial"], c["Weak"], c["Controls"]] for c in caf], [1.1, 2.5, 1.6, 0.85, 0.9, 0.85, 3.2], badge_cols=(2,))
    ssa = defaultdict(list)
    for c in controls:
        ssa[c["NcscPrinciple"]].append(c)
    ssa_titles = {"SSA 1": "Gain trust in your management devices", "SSA 2": "Protect your administration interfaces", "SSA 3": "Risk manage administration using tiers", "SSA 4": "Use privileged access management", "SSA 5": "Log and audit administration activities"}
    srows = []
    for k in ("SSA 1", "SSA 2", "SSA 3", "SSA 4", "SSA 5"):
        items = [c for c in ssa.get(k, []) if c["Result"] != "Not assessed"]
        if not items:
            rating = "Not assessed"
        elif any(c["Result"] == "Weak" and c["Severity"] in ("Critical", "High") for c in items):
            rating = "Weak"
        elif any(c["Result"] in ("Weak", "Partial") for c in items):
            rating = "Partial"
        else:
            rating = "Good"
        srows.append([k.replace("SSA ", ""), ssa_titles[k], rating, sum(1 for c in items if c["Result"] == "Good"), sum(1 for c in items if c["Result"] == "Partial"), sum(1 for c in items if c["Result"] == "Weak")])
    b.p("NCSC secure system administration principles:")
    b.table(["#", "Principle", "Position", "Good", "Partial", "Weak"], srows, [0.5, 4.5, 1.4, 0.8, 0.8, 0.8], badge_cols=(2,))

    # ---- 11 HLD alignment --------------------------------------------------------------------
    b.h("11. Alignment to the PIM Phase 2 HLD", 1)
    b.p("Status of each HLD current-state finding in this tenant. Not assessed means the evidence is not in the export, not that the finding is closed.")
    b.table(["Finding", "Title", "Status", "Measure"], [[h["Finding"], h["Title"], h["Status"], h["Measure"]] for h in hld], [1.0, 2.2, 1.3, 5.8], badge_cols=(2,))
    b.p("Maturity indicators against the HLD baseline (HLD 6.7):")
    b.table(["Capability", "HLD baseline", "Phase 2 target", "Measured indicator"], [[m["Capability"], m["HldBaseline"], m["Phase2Target"], m["Indicator"]] for m in maturity], [2.4, 1.5, 2.5, 4.0])

    # ---- 12 Plan -----------------------------------------------------------------------------
    b.h("12. Recommended Plan", 1)
    b.p("Actions are grouped by when they should be done, based on risk and dependency. Owners are to be agreed with the partner and the EIS operating teams.")
    by_id = {c["Id"]: c for c in controls}
    for title, ids in PHASES:
        items = [by_id[i] for i in ids if i in by_id and by_id[i]["Result"] in ("Weak", "Partial")]
        if not items:
            continue
        b.h(f"12.{PHASES.index((title, ids)) + 1} {title}", 2)
        b.table(["Ref", "Action", "Addresses", "Owner"], [[c["Id"], c["RecommendedAction"], c["Control"], "TBC"] for c in items], [0.8, 5.4, 3.0, 0.8])
    b.p("Progress can be measured by re-running the audit and comparing the results with this baseline (Compare-PimPhase2Assessment.ps1).", size=9.5)

    # ---- Appendices --------------------------------------------------------------------------
    b.page_break()
    b.h("Appendix A - All Control Checks", 1)
    b.table(["Ref", "Control", "Result", "Severity", "CAF", "Finding"],
            [[c["Id"], c["Control"], c["Result"], c["Severity"], c["CafOutcome"], c["Finding"]] for c in controls], [0.8, 2.6, 1.0, 0.9, 0.9, 4.6], badge_cols=(2, 3), size=8)
    b.h("Appendix B - TL1 and TL2 Roles in Use", 1)
    cat_rows = [c for c in catalogue if c["TrustLevel"] in ("TL1", "TL2") and c["Plane"] == "Entra"]
    b.table(["Level", "Role", "Classification", "People", "Direct", "Via group", "Permanent"],
            [[c["TrustLevel"], c["Role"], c["Classification"], c["People"], c["DirectPaths"], c["GroupPaths"], c["StandingPaths"]] for c in cat_rows],
            [0.8, 4, 1.4, 0.9, 0.9, 1.0, 1.1], badge_cols=(0,), size=8)
    b.p("Classification: HLD = set in the HLD; Proposed = proposed for agreement; Open = decision outstanding (OD-02). Azure roles are classified by role and scope; see P2_RoleCatalogue.csv.", size=8.5, color="5C6770")
    b.h("Appendix C - Evidence Files", 1)
    b.table(["File", "Content"], [
        ["Privileged_Access_Audit_<tenant>.html", "Interactive version of this report"],
        ["P2_Controls.csv", "Every control check with finding, why it matters, recommended action and NCSC references"],
        ["P2_AccessPaths.csv", "Every route from an identity to a role (named)"],
        ["P2_PrivilegedIdentities.csv", "Every admin identity with account type, organisation, linked everyday account and issues (named)"],
        ["P2_SelfApproval.csv", "People who can approve their own elevation (named)"],
        ["P2_MigrationWorksheet.csv", "One row per privileged route with proposed action, for the rebaseline workshops"],
        ["P2_PimCoverage.csv / P2_OversubscribedRoles.csv", "PIM usage per role and role populations against guides"],
        ["P2_ExceptionCandidates.csv", "Access that will need an exception record"],
        ["Azure\\Raw\\*.csv", "Azure role assignments, Lighthouse, Key Vault and storage access, coverage gaps"],
    ], [4, 6])
    b.h("Appendix D - Method and Limits", 1)
    b.bullet("Evidence was collected read-only with Microsoft Graph and Azure Resource Manager using a delegated reviewer account.")
    b.bullet("Trust levels come from the role catalogue agreed for PIM Phase 2; entries marked Proposed or Open are drafts until agreed.")
    b.bullet("Admin accounts are linked to everyday accounts by name and naming convention; where no everyday account appears in the evidence the link is marked for confirmation.")
    b.bullet("Microsoft guidance used: fewer than five Global Administrators; fewer than ten privileged role assignments; no more than three subscription owners.")
    if gaps:
        lines = []
        if any(g["Scope"] == "/" for g in gaps):
            lines.append("role assignments at the tenant root could not be read")
        nv = sum(1 for g in gaps if g["Operation"] == "Resources not visible")
        if nv:
            lines.append(f"resources inside {nv} subscriptions could not be listed, so resource-level checks are incomplete")
        al = sum(1 for g in gaps if "Activity" in g["Operation"])
        if al:
            lines.append(f"the Activity Log could not be read for {al} subscriptions")
        b.bullet("Azure coverage limits in this run: " + "; ".join(lines) + ". Reader at the tenant root management group closes these gaps.")
    b.bullet("References: NCSC CAF 4.0 Principle B2 (ncsc.gov.uk/collection/cyber-assessment-framework); NCSC Secure system administration (ncsc.gov.uk/collection/secure-system-administration); Microsoft best practices for Entra roles and Azure RBAC (learn.microsoft.com).")

    normalise_order(b.doc)
    b.doc.save(args.out)
    return args.out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--phase2", required=True)
    ap.add_argument("--azure-raw")
    ap.add_argument("--template", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--tenant-name")
    ap.add_argument("--owner", default="Andy Kemp - PIM Subject Matter Expert")
    ap.add_argument("--version", default="0.1")
    ap.add_argument("--audience", default="EIS Architecture, Cyber Security, IDS Operations and Shared Partnership stakeholders")
    ap.add_argument("--classification", default="Client Confidential")
    print(build(ap.parse_args()))


if __name__ == "__main__":
    main()
